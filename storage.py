"""
Ablage der Aufnahmen:
- Bild in den Unterordner des Fertigungsauftrags  (…\Bilder Mikroskop\81678\…)
- Eintrag in die gemeinsame Excel-Liste im Hauptordner
- Offline-Puffer: Ist der Server nicht erreichbar oder Excel gesperrt, wird lokal
  zwischengespeichert (%LOCALAPPDATA%\MikroskopCapture\puffer) und später nachgetragen.
- Rückgängig, Befunde für den PDF-Bericht lesen, Log-Datei.
"""

import json
import logging
import logging.handlers
import os
import re
import threading
import time
import uuid
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    Workbook = None

log = logging.getLogger("mikroskop")

# Excel-Spalten (Name, Breite). Vorhandene Listen werden anhand der Überschriften zugeordnet,
# fehlende Spalten werden hinten angefügt.
EXCEL_COLUMNS = [("Datum", 11), ("Uhrzeit", 9), ("Fertigungsauftrag (Bel_Nr)", 16), ("Fall-Nr.", 16),
                 ("Bild-Nr.", 8), ("Platinen-Nr.", 22), ("Artikelbezeichnung", 26), ("Fehlerart", 22),
                 ("Bauteil", 32), ("Position", 14), ("Gehäuse", 14), ("Technologie", 12), ("Bauteiltyp", 16),
                 ("Schadensbeschreibung", 40), ("Vergrößerung", 14), ("Bearbeiter", 14), ("Dateiname", 44),
                 ("Link", 11), ("Vorschau", 22)]
RENAMED_COLUMNS = (("Serien-/Auftrags-Nr.", "Fertigungsauftrag (Bel_Nr)"),
                   ("Leiterplatte / Artikel-Nr.", "Platinen-Nr."))
THUMB_HEIGHT = 90
FA_COLUMN = "Fertigungsauftrag (Bel_Nr)"


# ---------------- Hilfsfunktionen ----------------
def local_data_dir():
    base = Path(os.environ.get("LOCALAPPDATA") or Path.home())
    d = base / "MikroskopCapture"
    d.mkdir(parents=True, exist_ok=True)
    return d


def setup_logging(app_dir):
    """Log-Datei neben dem Programm (Fallback: %LOCALAPPDATA%), rotierend 5 x 1 MB."""
    for d in (Path(app_dir) / "logs", local_data_dir() / "logs"):
        try:
            d.mkdir(parents=True, exist_ok=True)
            path = d / "mikroskop.log"
            handler = logging.handlers.RotatingFileHandler(path, maxBytes=1_000_000, backupCount=5,
                                                           encoding="utf-8")
            break
        except OSError:
            continue
    else:
        return None
    handler.setFormatter(logging.Formatter("%(asctime)s  %(levelname)-7s  %(message)s", "%Y-%m-%d %H:%M:%S"))
    log.setLevel(logging.INFO)
    log.handlers[:] = [handler]
    log.propagate = False
    return path


def safe_name(text):
    text = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", str(text).strip())
    return re.sub(r"\s+", "_", text)[:40]


def fa_folder(fa):
    return safe_name(fa) or "ohne_FA"


def hide_dir(path):
    try:
        import ctypes
        ctypes.windll.kernel32.SetFileAttributesW(str(path), 0x02)
    except Exception:
        pass


def _hidden_subdir(root, name):
    d = Path(root) / name
    if not d.exists():
        d.mkdir(parents=True)
        hide_dir(d)
    return d


# ---------------- Excel ----------------
def open_sheet(xlsx):
    """Öffnet/erstellt die Liste; liefert (wb, ws, {Überschrift: Spalte})."""
    if Workbook is None:
        raise RuntimeError("openpyxl ist nicht installiert (pip install openpyxl)")
    if xlsx.exists():
        wb = load_workbook(xlsx)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Schäden"
        ws.freeze_panes = "A2"
    cols = {c.value: c.column for c in ws[1] if c.value}
    for old, new in RENAMED_COLUMNS:            # ältere Listen umbenennen
        if old in cols and new not in cols:
            col = cols.pop(old)
            ws.cell(row=1, column=col, value=new)
            cols[new] = col
    for name, width in EXCEL_COLUMNS:
        if name not in cols:
            col = max(cols.values(), default=0) + 1
            cell = ws.cell(row=1, column=col, value=name)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="305496")
            ws.column_dimensions[get_column_letter(col)].width = width
            cols[name] = col
    return wb, ws, cols


def place_thumbs(ws, cols, thumb_dir):
    # openpyxl verliert beim Laden vorhandene Bilder -> alle Vorschaubilder neu einfügen
    ws._images = []
    if not thumb_dir.is_dir():
        return
    name_col, thumb_col = cols["Dateiname"], get_column_letter(cols["Vorschau"])
    for row in range(2, ws.max_row + 2):
        name = ws.cell(row=row, column=name_col).value
        t = thumb_dir / (Path(str(name)).stem + ".jpg") if name else None
        if t and t.exists():
            ws.add_image(XLImage(str(t)), f"{thumb_col}{row}")
            ws.row_dimensions[row].height = THUMB_HEIGHT * 0.78
        elif row in ws.row_dimensions:
            ws.row_dimensions[row].height = None


class Store:
    """rec (Aufnahme-Datensatz): id, root, fa_folder, filename, excel_name (None = kein Excel),
    values (Excel-Spalten), created"""

    def __init__(self, on_change=None):
        self.lock = threading.RLock()
        self.buffer_dir = local_data_dir() / "puffer"
        self.buffer_dir.mkdir(parents=True, exist_ok=True)
        self.on_change = on_change or (lambda n: None)

    # ---- Pfade ----
    @staticmethod
    def image_path(rec):
        return Path(rec["root"]) / rec["fa_folder"] / rec["filename"]

    @staticmethod
    def new_record(root, fa, filename, excel_name, values):
        return {"id": uuid.uuid4().hex, "root": str(root), "fa_folder": fa_folder(fa), "filename": filename,
                "excel_name": excel_name, "values": values, "created": time.time()}

    # ---- Speichern ----
    def save(self, rec, img_bytes, thumb_bytes):
        """Speichert Bild + Excel-Eintrag. Ergebnis: ("ok"|"excel_buffered"|"buffered", Grund)."""
        with self.lock:
            try:
                self._write_image(rec, img_bytes)
            except Exception as e:
                log.warning("Bild nicht speicherbar (%s) -> Puffer: %s", e, rec["filename"])
                self._buffer(rec, "image", img_bytes, thumb_bytes)
                return "buffered", str(e)
            log.info("Bild gespeichert: %s", self.image_path(rec))
            if not rec["excel_name"]:
                return "ok", None
            try:
                self._excel_append(rec, thumb_bytes)
            except Exception as e:
                reason = "Excel-Liste ist geöffnet/gesperrt" if isinstance(e, PermissionError) else str(e)
                log.warning("Excel-Eintrag nicht möglich (%s) -> Puffer: %s", reason, rec["filename"])
                self._buffer(rec, "excel", None, thumb_bytes)
                return "excel_buffered", reason
            log.info("Excel-Eintrag: %s", rec["values"].get("Dateiname"))
            return "ok", None

    def _write_image(self, rec, img_bytes):
        d = Path(rec["root"]) / rec["fa_folder"]
        d.mkdir(parents=True, exist_ok=True)
        base, ext = os.path.splitext(rec["filename"])
        path, n = d / rec["filename"], 2
        while path.exists():                     # Namenskonflikt -> _2, _3 …
            path = d / f"{base}_{n}{ext}"
            n += 1
        path.write_bytes(img_bytes)              # funktioniert auch mit Umlauten / UNC-Pfaden
        rec["filename"] = path.name
        rec["values"]["Dateiname"] = path.name

    def _excel_append(self, rec, thumb_bytes):
        root = Path(rec["root"])
        xlsx = root / rec["excel_name"]
        if thumb_bytes:
            (_hidden_subdir(root, ".thumbs") / (Path(rec["filename"]).stem + ".jpg")).write_bytes(thumb_bytes)
        wb, ws, cols = open_sheet(xlsx)
        r = ws.max_row + 1
        for name, v in rec["values"].items():
            if name in cols:
                cell = ws.cell(row=r, column=cols[name], value=v)
                cell.alignment = Alignment(vertical="top", wrap_text=(name == "Schadensbeschreibung"))
        link = ws.cell(row=r, column=cols["Link"], value="Bild öffnen")
        link.hyperlink = str(self.image_path(rec))
        link.font = Font(color="0563C1", underline="single")
        link.alignment = Alignment(vertical="top")
        place_thumbs(ws, cols, root / ".thumbs")
        wb.save(xlsx)

    # ---- Offline-Puffer ----
    def _buffer(self, rec, stage, img_bytes, thumb_bytes):
        rec["stage"] = stage
        b = self.buffer_dir
        if img_bytes is not None:
            (b / f"{rec['id']}.img").write_bytes(img_bytes)
        if thumb_bytes:
            (b / f"{rec['id']}.thumb").write_bytes(thumb_bytes)
        (b / f"{rec['id']}.json").write_text(json.dumps(rec, ensure_ascii=False, indent=1), encoding="utf-8")
        self.on_change(self.pending_count())

    def pending_count(self):
        return len(list(self.buffer_dir.glob("*.json")))

    def is_buffered(self, rec):
        return (self.buffer_dir / f"{rec['id']}.json").exists()

    def buffered_image(self, rec):
        p = self.buffer_dir / f"{rec['id']}.img"
        return p if p.exists() else None

    def _drop(self, rid):
        for ext in ("json", "img", "thumb"):
            (self.buffer_dir / f"{rid}.{ext}").unlink(missing_ok=True)

    def retry(self):
        """Puffer nachtragen. Liefert (erledigt, verbleibend, letzter Fehler)."""
        done, err = 0, None
        with self.lock:
            for j in sorted(self.buffer_dir.glob("*.json"), key=lambda p: p.stat().st_mtime):
                try:
                    rec = json.loads(j.read_text(encoding="utf-8"))
                    rid = rec["id"]
                    thumb = self.buffer_dir / f"{rid}.thumb"
                    thumb_bytes = thumb.read_bytes() if thumb.exists() else None
                    if rec.get("stage") == "image":
                        self._write_image(rec, (self.buffer_dir / f"{rid}.img").read_bytes())
                        rec["stage"] = "excel"
                        j.write_text(json.dumps(rec, ensure_ascii=False, indent=1), encoding="utf-8")
                        (self.buffer_dir / f"{rid}.img").unlink(missing_ok=True)
                        log.info("Puffer: Bild nachgetragen: %s", self.image_path(rec))
                    if rec.get("excel_name"):
                        self._excel_append(rec, thumb_bytes)
                        log.info("Puffer: Excel-Eintrag nachgetragen: %s", rec["filename"])
                    self._drop(rid)
                    done += 1
                except Exception as e:
                    err = "Excel-Liste ist geöffnet/gesperrt" if isinstance(e, PermissionError) else str(e)
                    log.info("Puffer: Nachtragen noch nicht möglich (%s)", err)
                    break          # Server/Excel noch nicht bereit -> später erneut
        n = self.pending_count()
        if done:
            self.on_change(n)
        return done, n, err

    # ---- Rückgängig ----
    def undo(self, rec):
        """Entfernt eine Aufnahme (Bild -> .papierkorb, Excel-Zeile löschen). Liefert Warntext oder None."""
        with self.lock:
            if self.is_buffered(rec) and rec.get("stage", "image") == "image":
                self._drop(rec["id"])
                self.on_change(self.pending_count())
                log.info("Rückgängig (aus Puffer): %s", rec["filename"])
                return None
            excel_pending = self.is_buffered(rec)
            if excel_pending:
                self._drop(rec["id"])
                self.on_change(self.pending_count())
            warn = None
            root, path = Path(rec["root"]), self.image_path(rec)
            try:
                if path.exists():   # nicht endgültig löschen, sondern in versteckten Papierkorb-Ordner
                    os.replace(path, _hidden_subdir(root, ".papierkorb") / path.name)
                (root / ".thumbs" / (path.stem + ".jpg")).unlink(missing_ok=True)
            except Exception as e:
                warn = f"Bild konnte nicht entfernt werden: {e}"
            if rec["excel_name"] and not excel_pending:
                try:
                    self.remove_excel_row(root / rec["excel_name"], rec["filename"])
                except PermissionError:
                    warn = "Excel-Liste ist geöffnet – Zeile bitte von Hand löschen."
                except Exception as e:
                    warn = f"Excel-Zeile konnte nicht entfernt werden: {e}"
            log.info("Rückgängig: %s%s", rec["filename"], f" (Warnung: {warn})" if warn else "")
            return warn

    def remove_excel_row(self, xlsx, filename):
        wb, ws, cols = open_sheet(xlsx)
        for row in range(ws.max_row, 1, -1):
            if ws.cell(row=row, column=cols["Dateiname"]).value == filename:
                ws.delete_rows(row)
                break
        place_thumbs(ws, cols, xlsx.parent / ".thumbs")
        wb.save(xlsx)

    # ---- Befunde eines Auftrags (für den PDF-Bericht) ----
    def findings(self, root, excel_name, fa):
        """Alle Excel-Zeilen zum Fertigungsauftrag als dicts (+ "_image": Pfad zum Bild oder None)."""
        root = Path(root)
        xlsx = root / excel_name
        if not xlsx.exists():
            return []
        wb = load_workbook(xlsx, read_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        header = [str(h) if h is not None else "" for h in next(rows, [])]
        for old, new in RENAMED_COLUMNS:
            header = [new if h == old else h for h in header]
        out = []
        for r in rows:
            d = {h: ("" if v is None else v) for h, v in zip(header, r) if h}
            if str(d.get(FA_COLUMN, "")).strip() != str(fa).strip():
                continue
            name = str(d.get("Dateiname", ""))
            img = None
            for cand in (root / fa_folder(fa) / name, root / name):
                if name and cand.exists():
                    img = cand
                    break
            d["_image"] = img
            out.append(d)
        wb.close()
        return out
