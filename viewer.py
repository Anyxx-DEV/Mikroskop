"""
Fenster "Befunde": die Datensätze aus dbo.SMD_Mikroskop_Befunde wie eine Excel-Liste ansehen.
Suche, Filter (Fertigungsauftrag, Fehlerart, Zeitraum), Sortieren per Klick auf die Spaltenüberschrift,
Bildvorschau, Bild/Ordner öffnen, PDF-Bericht zum Auftrag und Export nach Excel (nur auf Knopfdruck).
"""

import os
import threading
import tkinter as tk
from datetime import datetime, timedelta
from pathlib import Path
from tkinter import filedialog, ttk

import customtkinter as ctk
from PIL import Image

import sqldb
import storage
import ui
from icons import ctk_icon
from storage import log
from ui import (BG, CARD, FIELD, BORDER, FG, MUTED, SECONDARY, SECONDARY_HOVER, OK_GREEN, WARN, ERR,
                F, button, entry, option, mode_color)

# (Schlüssel, Überschrift, Breite)
COLUMNS = [("Erfasst", "Datum / Uhrzeit", 140), ("Bel_Nr", "Fertigungsauftrag", 150), ("Bild", "Bild", 70),
           ("Platinen_Nr", "Platinen-Nr.", 130), ("Artikelbezeichnung", "Bezeichnung", 190),
           ("Fehlerart", "Fehlerart", 170), ("Bauteil", "Bauteil", 230), ("Position", "Position", 100),
           ("Beschreibung", "Beschreibung", 280), ("Bearbeiter", "Bearbeiter", 110)]
PERIODS = {"Heute": 0, "7 Tage": 7, "30 Tage": 30, "12 Monate": 365, "Alle": None}
ALL = "Alle Fehlerarten"
LIMIT = 2000


def _fmt(row, key):
    v = row.get(key)
    if key == "Erfasst":
        return v.strftime("%d.%m.%Y  %H:%M") if isinstance(v, datetime) else str(v or "")
    if key == "Bild":
        return f"{row.get('Bild_Nr') or ''}"
    if key == "Bauteil":
        return "  ·  ".join(x for x in (row.get("Bauteil_ArtikelNr") or "", row.get("Bauteil_Bez") or "") if x)
    if key == "Beschreibung":
        return " ".join(str(v or "").split())       # Zeilenumbrüche für die Tabelle entfernen
    return "" if v is None else str(v)


class BefundeWindow(ctk.CTkToplevel):
    def __init__(self, app, fa=""):
        super().__init__(app.root)
        self.app = app
        self.title("Befunde – Schadensdokumentation")
        self.configure(fg_color=BG)
        self.geometry("1500x900")
        self.minsize(1100, 650)
        self.rows = []
        self._sort = ("Erfasst", True)
        self._seq = 0
        self._job = None
        self._preview = None
        self._build(fa)
        self.after(60, lambda: (self.state("zoomed"), self.lift(), self.focus_force()))
        self.reload()

    # ---------------- Aufbau ----------------
    def _build(self, fa):
        top = ctk.CTkFrame(self, fg_color="transparent")
        top.pack(fill="x", padx=20, pady=(16, 8))
        ctk.CTkLabel(top, text="", image=ctk_icon("table", 26, FG)).pack(side="left")
        ctk.CTkLabel(top, text="Befunde", text_color=FG, font=F(22, "bold", ui.FONT_DISPLAY)).pack(side="left", padx=10)
        self.count = ctk.CTkLabel(top, text="", text_color=MUTED, font=F(13))
        self.count.pack(side="left", padx=8)
        button(top, " Excel-Export", self.export_excel, icon="download", kind="ghost").pack(side="right")
        button(top, " Aktualisieren", self.reload, icon="refresh").pack(side="right", padx=8)

        # Filterleiste
        bar = ctk.CTkFrame(self, fg_color=CARD, corner_radius=12, border_width=1, border_color=BORDER)
        bar.pack(fill="x", padx=20, pady=(0, 10))
        inner = ctk.CTkFrame(bar, fg_color="transparent")
        inner.pack(fill="x", padx=14, pady=12)
        self.text_var = tk.StringVar()
        self.fa_var = tk.StringVar(value=fa)
        ctk.CTkLabel(inner, text="", image=ctk_icon("search", 18, MUTED)).pack(side="left")
        e = entry(inner, self.text_var, width=320,
                  placeholder_text="Suchen: FA, Platine, Bauteil, Position, Beschreibung …")
        e.pack(side="left", padx=(6, 14))
        e.bind("<KeyRelease>", lambda ev: self._debounced_reload(), add="+")
        ctk.CTkLabel(inner, text="Fertigungsauftrag", text_color=MUTED, font=F(12, "bold")).pack(side="left")
        e = entry(inner, self.fa_var, width=110)
        e.pack(side="left", padx=(6, 14))
        e.bind("<KeyRelease>", lambda ev: self._debounced_reload(), add="+")
        self.fehler = option(inner, [ALL] + list(self.app.cfg.get("fehlerarten", [])),
                             lambda _: self.reload(), width=200, dynamic_resizing=False)
        self.fehler.set(ALL)
        self.fehler.pack(side="left", padx=(0, 14))
        self.period = ctk.CTkSegmentedButton(
            inner, values=list(PERIODS), height=34, corner_radius=8, font=F(12, "bold"),
            fg_color=(SECONDARY[0], FIELD[1]), unselected_color=(SECONDARY[0], FIELD[1]),
            unselected_hover_color=SECONDARY_HOVER, selected_color=("#FFFFFF", "#475569"),
            selected_hover_color=("#FFFFFF", "#475569"), text_color=FG, command=lambda _: self.reload())
        self.period.set("Alle" if fa else "30 Tage")
        self.period.pack(side="left")
        button(inner, "Filter zurücksetzen", self.reset_filters, kind="ghost", height=34).pack(side="right")

        # Tabelle + Detailbereich
        body = ctk.CTkFrame(self, fg_color="transparent")
        body.pack(fill="both", expand=True, padx=20, pady=(0, 8))
        body.grid_columnconfigure(0, weight=1)
        body.grid_rowconfigure(0, weight=1)

        tbl = ctk.CTkFrame(body, fg_color=CARD, corner_radius=12, border_width=1, border_color=BORDER)
        tbl.grid(row=0, column=0, sticky="nsew")
        tbl.grid_columnconfigure(0, weight=1)
        tbl.grid_rowconfigure(0, weight=1)
        self._style()
        self.tree = ttk.Treeview(tbl, columns=[c for c, _, _ in COLUMNS], show="headings",
                                 style="Befunde.Treeview", selectmode="browse")
        for key, title, width in COLUMNS:
            self.tree.heading(key, text=title, anchor="w", command=lambda k=key: self.sort_by(k))
            self.tree.column(key, width=int(width * self._scale), minwidth=50, anchor="w",
                             stretch=key in ("Beschreibung", "Bauteil"))
        self.tree.grid(row=0, column=0, sticky="nsew", padx=(8, 0), pady=(8, 0))
        ys = ctk.CTkScrollbar(tbl, command=self.tree.yview)
        ys.grid(row=0, column=1, sticky="ns", pady=8)
        xs = ctk.CTkScrollbar(tbl, orientation="horizontal", command=self.tree.xview)
        xs.grid(row=1, column=0, sticky="ew", padx=8)
        self.tree.configure(yscrollcommand=ys.set, xscrollcommand=xs.set)
        self.tree.bind("<<TreeviewSelect>>", lambda e: self._show_selected())
        self.tree.bind("<Double-1>", lambda e: self.open_image())
        self.tree.bind("<Return>", lambda e: self.open_image())
        self.empty = ctk.CTkLabel(tbl, text="", text_color=MUTED, font=F(14), justify="center")

        # Detailbereich rechts
        side = ctk.CTkFrame(body, fg_color=CARD, corner_radius=12, border_width=1, border_color=BORDER, width=420)
        side.grid(row=0, column=1, sticky="ns", padx=(12, 0))
        side.grid_propagate(False)
        # leeres Platzhalterbild: CTkLabel entfernt ein Bild mit image=None nicht zuverlässig
        blank = Image.new("RGBA", (1, 1), (0, 0, 0, 0))
        self._blank = ctk.CTkImage(light_image=blank, dark_image=blank, size=(1, 1))
        self.preview = ctk.CTkLabel(side, text="Befund auswählen", text_color=MUTED, font=F(13),
                                    fg_color=FIELD, corner_radius=8, height=240, image=self._blank)
        self.preview.pack(fill="x", padx=14, pady=(14, 10))
        self.detail = ctk.CTkFrame(side, fg_color="transparent")
        self.detail.pack(fill="both", expand=True, padx=14)
        btns = ctk.CTkFrame(side, fg_color="transparent")
        btns.pack(fill="x", padx=14, pady=14, side="bottom")
        self.btn_img = button(btns, " Bild öffnen", self.open_image, icon="image", state="disabled")
        self.btn_img.pack(fill="x")
        self.btn_dir = button(btns, " Ordner öffnen", self.open_folder, icon="folder_open", kind="ghost",
                              state="disabled")
        self.btn_dir.pack(fill="x", pady=6)
        self.btn_pdf = button(btns, " PDF-Bericht zum Auftrag", self.pdf_report, icon="clipboard", kind="ghost",
                              state="disabled")
        self.btn_pdf.pack(fill="x")

        self.status = ctk.CTkLabel(self, text="", text_color=MUTED, font=F(12), anchor="w")
        self.status.pack(fill="x", padx=24, pady=(0, 12))

    def _style(self):
        self._scale = ctk.ScalingTracker.get_widget_scaling(self)
        st = ttk.Style(self)
        st.theme_use("clam")
        bg, fg, head = mode_color(CARD), mode_color(FG), mode_color(SECONDARY)
        sel = "#14532D" if ctk.get_appearance_mode() == "Dark" else "#D1FAE5"
        # Schriftgröße in Pixeln (negativ) und selbst skaliert - ttk wächst sonst nicht mit der Windows-Skalierung
        px = -int(14 * self._scale)
        st.configure("Befunde.Treeview", background=bg, fieldbackground=bg, foreground=fg, borderwidth=0,
                     rowheight=int(34 * self._scale), font=(ui.FONT_UI, px))
        st.configure("Befunde.Treeview.Heading", background=head, foreground=fg, relief="flat", borderwidth=0,
                     font=(ui.FONT_UI, px, "bold"), padding=(int(8 * self._scale), int(8 * self._scale)))
        st.map("Befunde.Treeview", background=[("selected", sel)], foreground=[("selected", fg)])
        st.map("Befunde.Treeview.Heading", background=[("active", mode_color(SECONDARY_HOVER))])
        st.layout("Befunde.Treeview", [("Treeview.treearea", {"sticky": "nswe"})])   # ohne Rahmen
        self._odd_bg = mode_color(FIELD)

    # ---------------- Daten ----------------
    def _filters(self):
        days = PERIODS[self.period.get()]
        date_from = None
        if days is not None:
            today = datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)
            date_from = today - timedelta(days=days)
        f = self.fehler.get()
        return dict(text=self.text_var.get().strip(), fa=self.fa_var.get().strip(),
                    fehlerart="" if f == ALL else f, date_from=date_from, date_to=None)

    def _debounced_reload(self):
        if self._job:
            self.after_cancel(self._job)
        self._job = self.after(400, self.reload)

    def reset_filters(self):
        self.text_var.set("")
        self.fa_var.set("")
        self.fehler.set(ALL)
        self.period.set("30 Tage")
        self.reload()

    def reload(self):
        self._job = None
        cfg = self.app._sql_cfg()
        if not self.app._sql_ready():
            self._show_message("Datenbank nicht eingerichtet.\nIm Hauptfenster: Reiter „Einstellungen“ → „Datenbank“.")
            return
        self._seq += 1
        seq, flt = self._seq, self._filters()
        self.status.configure(text="Lade Befunde…", text_color=MUTED)

        def work():
            try:
                rows, total = sqldb.query_befunde(cfg, limit=LIMIT, **flt)
                err = None
            except Exception as e:
                rows, total, err = [], 0, storage._short_err(e)
            self.after(0, lambda: self._loaded(seq, rows, total, err))

        threading.Thread(target=work, daemon=True).start()

    def _loaded(self, seq, rows, total, err):
        if seq != self._seq or not self.winfo_exists():
            return
        if err:
            log.error("Befunde-Liste: %s", err)
            hint = ("\n\nDie Tabelle SMD_Mikroskop_Befunde fehlt vermutlich – im Hauptfenster unter "
                    "„Einstellungen“ → „Datenbank“ anlegen." if "SMD_Mikroskop_Befunde" in err else "")
            self._show_message(f"Befunde konnten nicht geladen werden:\n{err}{hint}")
            self.status.configure(text=err, text_color=ERR)
            return
        self.rows = rows
        self._apply_sort()
        if total > len(rows):
            self.count.configure(text=f"{len(rows)} von {total} Befunden (neueste zuerst)")
            self.status.configure(text="Es werden die neuesten Befunde angezeigt – Filter verfeinern, um ältere zu "
                                       "finden.", text_color=WARN)
        else:
            self.count.configure(text=f"{total} Befund{'e' if total != 1 else ''}")
            self.status.configure(text="Doppelklick oder Enter öffnet das Bild · Klick auf eine Spaltenüberschrift "
                                       "sortiert", text_color=MUTED)

    def _show_message(self, text):
        self.tree.delete(*self.tree.get_children())
        self.rows = []
        self.count.configure(text="")
        self.empty.configure(text=text)
        self.empty.place(relx=0.5, rely=0.45, anchor="center")
        self._clear_detail()

    def _fill(self):
        self.empty.place_forget()
        self.tree.delete(*self.tree.get_children())
        self.tree.tag_configure("odd", background=self._odd_bg)
        for i, r in enumerate(self.rows):
            self.tree.insert("", "end", iid=str(i), values=[_fmt(r, k) for k, _, _ in COLUMNS],
                             tags=("odd",) if i % 2 else ())
        if not self.rows:
            self.empty.configure(text="Keine Befunde für diese Filter.")
            self.empty.place(relx=0.5, rely=0.45, anchor="center")
        self._clear_detail()

    # ---------------- Sortieren ----------------
    def sort_by(self, key):
        k, desc = self._sort
        self._sort = (key, not desc if k == key else False)
        self._apply_sort()

    def _apply_sort(self):
        key, desc = self._sort
        if key == "Erfasst":
            sort_key = lambda r: r.get("Erfasst") or datetime.min
        elif key == "Bild":
            sort_key = lambda r: (r.get("Bel_Nr") or "", r.get("Bild_Nr") or 0)
        else:
            sort_key = lambda r: _fmt(r, key).lower()
        self.rows.sort(key=sort_key, reverse=desc)
        for k, title, _ in COLUMNS:
            arrow = (" ▼" if desc else " ▲") if k == key else ""
            self.tree.heading(k, text=title + arrow)
        self._fill()

    # ---------------- Auswahl / Details ----------------
    def selected(self):
        sel = self.tree.selection()
        return self.rows[int(sel[0])] if sel else None

    def _clear_detail(self):
        for w in self.detail.winfo_children():
            w.destroy()
        self.preview.configure(image=self._blank, text="Befund auswählen")
        self._preview = None
        for b in (self.btn_img, self.btn_dir, self.btn_pdf):
            b.configure(state="disabled")
        self.btn_pdf.configure(text=" PDF-Bericht zum Auftrag")

    def _show_selected(self):
        r = self.selected()
        if not r:
            return
        for w in self.detail.winfo_children():
            w.destroy()
        rows = [("Erfasst", _fmt(r, "Erfasst")), ("Fertigungsauftrag", r.get("Bel_Nr")),
                ("Fall / Bild", f"{r.get('Fall_Nr') or ''} / {r.get('Bild_Nr') or ''}"),
                ("Platine", f"{r.get('Platinen_Nr') or ''}  {r.get('Artikelbezeichnung') or ''}".strip()),
                ("Fehlerart", r.get("Fehlerart")), ("Bauteil", _fmt(r, "Bauteil")),
                ("Position", r.get("Position")),
                ("Details", ", ".join(x for x in (r.get("Gehaeuse"), r.get("Technologie"), r.get("Bauteiltyp")) if x)),
                ("Vergrößerung", r.get("Vergroesserung")), ("Bearbeiter", r.get("Bearbeiter")),
                ("Arbeitsplatz", r.get("Arbeitsplatz")), ("Beschreibung", r.get("Beschreibung"))]
        for k, v in rows:
            if not v:
                continue
            row = ctk.CTkFrame(self.detail, fg_color="transparent")
            row.pack(fill="x", pady=2)
            ctk.CTkLabel(row, text=k, text_color=MUTED, font=F(12), width=118, anchor="nw").pack(side="left", anchor="n")
            ctk.CTkLabel(row, text=str(v), text_color=FG, font=F(12, "bold"), anchor="w", justify="left",
                         wraplength=260).pack(side="left", fill="x", expand=True)
        has_img = bool(r.get("Bildpfad"))
        self.btn_img.configure(state="normal" if has_img else "disabled")
        self.btn_dir.configure(state="normal" if has_img else "disabled")
        self.btn_pdf.configure(state="normal", text=f" PDF-Bericht FA {r.get('Bel_Nr')}")
        self._load_preview(r)

    def _load_preview(self, r):
        path = r.get("Bildpfad")
        self.preview.configure(image=self._blank, text="Lade Vorschau…")
        if not path:
            self.preview.configure(text="Kein Bild")
            return
        sel_id = r.get("ID")

        def work():
            try:
                src = storage.thumb_path(path) or Path(path)
                with Image.open(src) as im:
                    im = im.convert("RGB")
                    im.thumbnail((390, 260))
                    img = im.copy()
                err = None
            except Exception:
                img, err = None, "Bild nicht gefunden"
            self.after(0, lambda: self._set_preview(sel_id, img, err))

        threading.Thread(target=work, daemon=True).start()

    def _set_preview(self, sel_id, img, err):
        r = self.selected()
        if not r or r.get("ID") != sel_id or not self.winfo_exists():
            return
        if img is None:
            self.preview.configure(image=self._blank, text=err)
            return
        self._preview = ctk.CTkImage(light_image=img, dark_image=img, size=img.size)
        self.preview.configure(image=self._preview, text="")

    # ---------------- Aktionen ----------------
    def open_image(self):
        r = self.selected()
        if r and r.get("Bildpfad"):
            try:
                os.startfile(r["Bildpfad"])
            except OSError as e:
                self.status.configure(text=f"Bild nicht erreichbar: {e}", text_color=ERR)

    def open_folder(self):
        r = self.selected()
        if r and r.get("Bildpfad"):
            try:
                os.startfile(str(Path(r["Bildpfad"]).parent))
            except OSError as e:
                self.status.configure(text=f"Ordner nicht erreichbar: {e}", text_color=ERR)

    def pdf_report(self):
        r = self.selected()
        if not r:
            return
        self.status.configure(text=f"Erstelle PDF-Bericht für FA {r['Bel_Nr']}…", text_color=MUTED)
        self.app.make_pdf_report(r["Bel_Nr"], on_done=lambda ok, res: self.winfo_exists() and self.status.configure(
            text=f"PDF-Bericht erstellt: {Path(res).name}" if ok else f"PDF-Bericht: {res}",
            text_color=OK_GREEN if ok else ERR))

    def export_excel(self):
        """Aktuell angezeigte Liste als Excel-Datei speichern (nur bei Bedarf, z. B. zum Weitergeben)."""
        if not self.rows:
            self.status.configure(text="Keine Befunde zum Exportieren.", text_color=WARN)
            return
        name = f"Befunde_{datetime.now():%Y-%m-%d}.xlsx"
        path = filedialog.asksaveasfilename(parent=self, title="Befunde als Excel speichern", initialfile=name,
                                            defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")])
        if not path:
            return
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter
            wb = Workbook()
            ws = wb.active
            ws.title = "Befunde"
            heads = [t for _, t, _ in COLUMNS] + ["Gehäuse", "Technologie", "Bauteiltyp", "Vergrößerung",
                                                 "Arbeitsplatz", "Bild"]
            ws.append(heads)
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="333333")
            for r in self.rows:
                vals = [_fmt(r, k) for k, _, _ in COLUMNS] + [r.get("Gehaeuse"), r.get("Technologie"),
                                                             r.get("Bauteiltyp"), r.get("Vergroesserung"),
                                                             r.get("Arbeitsplatz"), "Bild öffnen"]
                ws.append(vals)
                cell = ws.cell(row=ws.max_row, column=len(heads))
                if r.get("Bildpfad"):
                    cell.hyperlink = r["Bildpfad"]
                    cell.font = Font(color="0563C1", underline="single")
            for i, (_, _, w) in enumerate(COLUMNS, 1):
                ws.column_dimensions[get_column_letter(i)].width = max(10, w // 7)
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    c.alignment = Alignment(vertical="top")
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            wb.save(path)
            log.info("Befunde exportiert: %s (%s Zeilen)", path, len(self.rows))
            self.status.configure(text=f"Exportiert: {path}", text_color=OK_GREEN)
            os.startfile(path)
        except Exception as e:
            log.error("Excel-Export fehlgeschlagen: %s", e)
            self.status.configure(text=f"Export fehlgeschlagen: {e}", text_color=ERR)
